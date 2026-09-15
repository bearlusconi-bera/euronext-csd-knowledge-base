"""Build non-destructive structured derivatives from the retained audit evidence.

Run with the existing local Python environment (openpyxl and lxml are available).
No network access, source replacement or automatic admission is performed.
"""
import hashlib
import json
import re
import stat
import zipfile
from pathlib import Path, PurePosixPath

from lxml import etree, html
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "implementation/2026-09-13"
AUDIT = ROOT / "audits/2026-09-13"
CAT = {d["id"]: d for d in json.loads((ROOT / "catalogue.json").read_text())}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(name, value):
    path = OUT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def matching():
    """Transcription visually checked against UDFS PDF 269 and 270, all 3 diagrams."""
    source = CAT["aa3d5a3b94c9"]
    common = ["Payment Type", "Securities Movement Type", "ISIN Code", "Trade Date",
              "Settlement Quantity", "Intended Settlement Date", "Delivering Party BIC",
              "Receiving Party BIC", "CSD of the Delivering Party", "CSD of the Receiving Party"]
    cash = ["Currency", "Settlement Amount", "Credit/Debit"]
    additional = ["Opt-out ISO transaction condition indicator", "CUM/EX Indicator"]
    optional = ["Common Trade Reference", "Client of delivering CSD participant",
                "Client of receiving CSD participant", "Securities account of the delivering party",
                "Securities account of the receiving party"]
    rows = []
    for name in common + cash:
        rows.append({"field": name, "diagram": 55, "pdf_page": 269,
                     "DVP/DWP": "mandatory", "FOP": "mandatory" if name in common else "n/a"})
    for name in additional + cash:
        rows.append({"field": name, "diagram": 56, "pdf_page": 270,
                     "DVP/DWP": "additional" if name in additional else "n/a", "FOP": "additional"})
    for name in optional:
        rows.append({"field": name, "diagram": 57, "pdf_page": 270,
                     "DVP/DWP": "optional", "FOP": "optional"})
    result = {
        "source_id": source["id"], "source_sha256": source["sha256"], "source_url": source["url"],
        "release": "R2026.JUN", "verified_as_of": "2026-09-13", "actual_content_language": "en",
        "review": "All rows, diagram notes and footnote 194 visually inspected; PDF 271 tolerance narrative read.",
        "scope": "Functional matching-field diagrams, not production XML paths, XSD validation or complete message usage rules.",
        "transaction_headers_verbatim": ["DVP/DWP", "FOP"],
        "rows": rows,
        "conditions": [
            "Mandatory values match, except opposing Credit/Debit and Securities Movement Type values; settlement amount may use the stated tolerance. The paragraph mentions DVP/PFOD; the diagram header itself reads DVP/DWP. Do not silently normalise these labels into a certified schema mapping.",
            "Additional: a value supplied on either side must also be supplied and match on the other side; blank/blank matches. Credit/Debit matches opposite values.",
            "Optional: one filled and one blank can match; if both are filled they must match.",
            "Footnote 194: upper- and lower-case letters are considered different when comparing values. Do not case-normalise identifiers before comparing.",
            "Diagram 56 note 1: CUM/EX matching considers only ExCoupon and CumCoupon; other values are considered blank.",
            "Diagram 56 note 2: Currency, Settlement Amount and Credit/Debit are additional FOP fields to reduce mismatching risk for non-T2S-currency cash legs submitted as FOP (Payment Flag FREE) with CoSD used to ensure DVP.",
            "Diagram 57 note: client fields match BICs or proprietary codes. Proprietary code matching uses Identification, Issuer and Scheme Name. A BIC does not match a proprietary code.",
            "PDF 270-271: EUR amount tolerance is EUR 2 for cash countervalue <= EUR 100,000 and EUR 25 above EUR 100,000. Currency-specific configuration applies; do not generalise EUR bands to other currencies.",
            "PDF 271: among candidates choose the smallest amount difference, then closest entry time if amounts are the same; the deliverer's amount becomes the matched settlement amount."
        ],
        "counts": {"mandatory_diagram_rows": 13, "additional_diagram_rows": 5, "optional_diagram_rows": 5},
        "images": ["audits/2026-09-13/evidence/aa3d5a3b94c9-p269.png",
                   "implementation/2026-09-13/evidence/t2s-p270.png"],
        "production_schema_validated": False,
    }
    assert len(rows) == 23
    save("structured/t2s-matching-fields.json", result)


def workbook():
    d = CAT["ac385b0f3106"]
    wb = load_workbook(ROOT / d["local_path"], data_only=True, read_only=False)
    sheets = []
    for ws in wb.worksheets:
        cells = [{"address": c.coordinate, "value": c.value} for row in ws for c in row if c.value is not None]
        sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column,
                       "merged_ranges": [str(r) for r in ws.merged_cells.ranges], "cells": cells})
    save("structured/european-offering-workbook.json", {
        "source_id": d["id"], "source_sha256": d["sha256"], "source_url": d["url"],
        "scope": "European Offering listing snapshot; not universal Milan settlement eligibility",
        "snapshot_date": "2026-09-11", "data_through": "2026-09-10",
        "designated_and_alternative_place_columns_apply_from": "2026-09-21",
        "column_date_evidence": "Introduction B5:B8; preserve each sheet's header cells and all alternative columns",
        "default_retrieval": False, "review_scope": "Introduction, column headers and audit sample only; rows mechanically extracted, not individually validated",
        "sheets": sheets,
    })


def events():
    meta = json.loads((OUT / "evidence/ecb-status-2026.json").read_text())
    doc = html.fromstring((ROOT / meta["local_path"]).read_bytes())
    text = doc.text_content()
    matches = list(re.finditer(r"(?P<d>\d{2}/\d{2}/2026)\s*(?P<t>\d{2}:\d{2}:\d{2})", text))
    records = []
    for i, m in enumerate(matches):
        day, month, year = m["d"].split("/")
        body = text[m.end():matches[i+1].start() if i+1 < len(matches) else len(text)].strip()
        body = re.sub(r"English\s*$", "", body).strip()
        if f"{year}-{month}-{day}" != "2026-09-08":
            continue
        records.append({"id": "t2s-20260908-" + m["t"].replace(":", ""),
                        "business_date": "2026-09-08", "published_time_as_displayed": m["t"],
                        "timestamp_timezone": None, "timezone_note": "Preserve displayed ECB timestamp; UTC conversion not verified",
                        "original_text": body, "source_id": "ecb-status-2026", "source_sha256": meta["sha256"],
                        "source_url": meta["url"], "reviewed": True})
    records.sort(key=lambda x: x["published_time_as_displayed"])
    by_time = {e["published_time_as_displayed"]: e for e in records}
    by_time["15:30:00"].update(event_type="announced_schedule_change", currency="EUR", event="IDVP",
                              original_time="16:00", announced_time="17:00", actual_completion_time=None,
                              baseline_time_convention="CET as stated in UDFS; no UTC conversion",
                              scope="EUR DvP cutoff on this business date only",
                              dependent_events=["RMIC and OCSW 15 minutes after IDVP closure"],
                              supersedes_event_ids=[])
    by_time["18:10:00"].update(event_type="incident_update", currency=None, event="ICBO/IFOP",
                              scope="Operating day blocked; IFOP had not started", supersedes_event_ids=[])
    by_time["18:55:00"].update(event_type="incident_resolution", currency=None, event="IFOP",
                              actual_completion_time="18:39", scope="ICBO incident resolved; IFOP closure reported",
                              supersedes_event_ids=[by_time["18:10:00"]["id"]])
    save("structured/t2s-events.json", {
        "coverage": {"2026-09-08": "All displayed updates for this date retained and read"},
        "verified_as_of": "2026-09-13", "events": records,
        "policy": "Apply only to matching business date/currency/event. Later general normal-status messages do not erase earlier historical exceptions. No evidence for unreviewed dates."})


def entity():
    raw = json.loads((OUT / "sources/gleif-interbolsa.json").read_text())
    a = raw["data"]["attributes"]
    lei = a["lei"]
    expanded = "".join(str(ord(c)-55) if c.isalpha() else c for c in lei)
    assert len(lei) == 20 and int(expanded) % 97 == 1
    save("structured/interbolsa-identity.json", {
        "entity": a["entity"]["legalName"]["name"], "normalised_lei": lei,
        "normalised_lei_source": "https://api.gleif.org/api/v1/lei-records/" + lei,
        "source_sha256": digest(OUT / "sources/gleif-interbolsa.json"),
        "json_pointers": ["/data/attributes/lei", "/data/attributes/entity/legalName", "/data/attributes/registration"],
        "entity_status": a["entity"]["status"], "registration_status": a["registration"]["status"],
        "last_update_date": a["registration"]["lastUpdateDate"], "retrieved_as_of": "2026-09-13",
        "raw_esma_cell": {"source_id": "535f6a23062a", "sheet": "1- CSD authorisations", "cell": "C25", "value": "529900LG70TCA", "usable_for_identity_join": False},
        "checksum_valid": True, "note": "Normalisation is independently sourced; original ESMA workbook and cell extract remain unchanged."
    })


def archives():
    """Inspect every archived container, stage members by hash, never execute contents."""
    results = []
    for d in CAT.values():
        if d.get("file_type") != "zip" or not d.get("local_path"):
            continue
        p = ROOT / d["local_path"]
        record = {"source_id": d["id"], "source_sha256": digest(p), "members": [], "default_retrieval": False}
        with zipfile.ZipFile(p) as z:
            record["detected_type"] = "xlsx" if "xl/workbook.xml" in z.namelist() else "zip"
            total = 0
            for info in z.infolist():
                if info.is_dir():
                    continue
                name = PurePosixPath(info.filename.replace("\\", "/"))
                mode = info.external_attr >> 16
                blocked = (name.is_absolute() or ".." in name.parts or stat.S_ISLNK(mode) or
                           info.file_size > 32_000_000 or total + info.file_size > 160_000_000 or
                           info.flag_bits & 1 or (info.compress_size and info.file_size/info.compress_size > 500))
                m = {"name": info.filename, "bytes": info.file_size, "admitted": False}
                if blocked:
                    m["status"] = "blocked-unsafe-or-oversized-member"
                    record["members"].append(m)
                    continue
                total += info.file_size
                data = z.read(info)
                sha = hashlib.sha256(data).hexdigest()
                suffix = name.suffix.lower()
                target = OUT / "staged-archive-members" / (sha + suffix)
                target.parent.mkdir(parents=True, exist_ok=True)
                if not target.exists():
                    target.write_bytes(data)
                m.update(sha256=sha, staged_path=str(target.relative_to(ROOT)), status="staged-not-semantically-reviewed")
                if suffix in [".xml", ".xsd", ".rels"]:
                    try:
                        node = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
                        m["xml_well_formed"] = True
                        refs = node.xpath('//*[local-name()="import" or local-name()="include" or local-name()="redefine"]/@schemaLocation')
                        m["schema_references"] = refs
                        m["schema_validated"] = False
                        m["unresolved_schema_references"] = []
                        for ref in refs:
                            import posixpath
                            target_name = posixpath.normpath(str(name.parent / ref))
                            if "://" in ref or target_name not in z.namelist():
                                m["unresolved_schema_references"].append(ref)
                    except etree.XMLSyntaxError as e:
                        m["xml_well_formed"] = False
                        m["parse_error"] = str(e)
                record["members"].append(m)
        results.append(record)
    save("structured/archive-members.json", results)


if __name__ == "__main__":
    matching()
    workbook()
    events()
    entity()
    archives()
    print("Structured repairs generated; no original changed or whole document admitted.")
